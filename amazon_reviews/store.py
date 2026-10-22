"""Deduplicating review store, crash-safe checkpointing, and exports."""

from __future__ import annotations

import csv
import json
import logging
from pathlib import Path

from .models import Review

log = logging.getLogger("amazon_reviews.store")


def export_combined(
    stores: list["ReviewStore"], out_dir: Path, formats: list[str], stem: str = "all_products"
) -> list[Path]:
    """Write every product's reviews into one merged file.

    Each row already carries its `asin`, so the merged sheet can be filtered or
    pivoted per product. Written alongside the per-product files, never instead
    of them - one product failing should not cost you the rest.
    """
    stores = [s for s in stores if len(s)]
    if not stores:
        return []

    merged = ReviewStore(stem, out_dir, cross_product=True)
    for store in stores:
        merged.add_many(store.reviews)
    # Sum the per-product totals so coverage still means something.
    totals = [s.total_reviews_reported for s in stores if s.total_reviews_reported]
    if totals:
        merged.total_reviews_reported = sum(totals)
    for store in stores:
        merged.capped_segments |= {f"{store.asin}:{c}" for c in store.capped_segments}

    written = merged.export(formats)
    merged.clear_checkpoint()
    return written


class ReviewStore:
    """Collects reviews across segments, keyed by review id.

    Segments overlap heavily by design, so the same review arrives many times.
    First write wins, which keeps `found_via` pointing at the segment that
    actually discovered it - handy for auditing whether a coverage level is
    pulling its weight.
    """

    def __init__(
        self,
        asin: str,
        out_dir: Path,
        star_filter: set[int] | None = None,
        cross_product: bool = False,
    ):
        # Within one product, the review id is the identity. Across products it
        # is not: Amazon shows the same review under every variant ASIN, so a
        # merged file must keep one row per (product, review) or it silently
        # drops rows and attributes them to whichever product was seen first.
        self.cross_product = cross_product
        self.asin = asin
        self.out_dir = Path(out_dir)
        self.out_dir.mkdir(parents=True, exist_ok=True)
        self._reviews: dict[str, Review] = {}
        self.completed_segments: set[str] = set()
        self.total_reviews_reported: int | None = None
        self.total_ratings_reported: int | None = None
        # Which star ratings the user asked for; None means everything.
        self.star_filter = set(star_filter) if star_filter else None
        # Segments that stopped at Amazon's cap rather than running out. These
        # are the pools worth attacking with keyword segments.
        self.capped_segments: set[str] = set()
        # True when review pages demanded a sign-in, so anything collected came
        # only from the product page and is a fraction of what exists.
        self.login_blocked: bool = False

    # -- collection -------------------------------------------------------

    def _wanted(self, review: Review) -> bool:
        if self.star_filter is None:
            return True
        if review.rating is None:
            return False
        return int(review.rating) in self.star_filter

    def _key(self, review: Review):
        return (review.asin, review.review_id) if self.cross_product else review.review_id

    def add_many(self, reviews: list[Review]) -> int:
        """Add reviews, returning how many were new.

        Reviews outside the requested star range are dropped here rather than
        at export, so the running totals the UI shows always mean the same
        thing as the final file.
        """
        new = 0
        for review in reviews:
            key = self._key(review)
            if key not in self._reviews and self._wanted(review):
                self._reviews[key] = review
                new += 1
        return new

    def __len__(self) -> int:
        return len(self._reviews)

    @property
    def reviews(self) -> list[Review]:
        """All reviews, newest first; undated ones sort last."""
        return sorted(
            self._reviews.values(),
            key=lambda r: (r.review_date or "0000-00-00"),
            reverse=True,
        )

    @property
    def coverage_pct(self) -> float | None:
        if not self.total_reviews_reported:
            return None
        return 100.0 * len(self._reviews) / self.total_reviews_reported

    def summary(self) -> dict:
        by_star: dict[str, int] = {}
        for review in self._reviews.values():
            key = f"{int(review.rating)}_star" if review.rating else "unrated"
            by_star[key] = by_star.get(key, 0) + 1

        verified = sum(1 for r in self._reviews.values() if r.verified_purchase)
        rated = [r.rating for r in self._reviews.values() if r.rating]
        return {
            "asin": self.asin,
            "reviews_downloaded": len(self._reviews),
            "reviews_reported_by_amazon": self.total_reviews_reported,
            "ratings_reported_by_amazon": self.total_ratings_reported,
            "coverage_pct": round(self.coverage_pct, 1) if self.coverage_pct else None,
            "average_rating": round(sum(rated) / len(rated), 2) if rated else None,
            "verified_purchases": verified,
            "with_images": sum(1 for r in self._reviews.values() if r.images),
            "by_star": dict(sorted(by_star.items(), reverse=True)),
            "segments_completed": len(self.completed_segments),
            # A non-empty list means Amazon truncated those pools; the numbers
            # for those ratings are floors, not totals.
            "capped_segments": sorted(self.capped_segments),
            "login_blocked": self.login_blocked,
        }

    # -- checkpointing ----------------------------------------------------

    @property
    def checkpoint_path(self) -> Path:
        return self.out_dir / f".checkpoint_{self.asin}.json"

    def save_checkpoint(self) -> None:
        """Write progress so an interrupted run can pick up where it stopped."""
        payload = {
            "asin": self.asin,
            "completed_segments": sorted(self.completed_segments),
            "total_reviews_reported": self.total_reviews_reported,
            "total_ratings_reported": self.total_ratings_reported,
            "reviews": [r.to_row() | {"images": r.images} for r in self._reviews.values()],
        }
        tmp = self.checkpoint_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        tmp.replace(self.checkpoint_path)  # atomic: never leave a half-written file

    def load_checkpoint(self) -> bool:
        if not self.checkpoint_path.exists():
            return False
        try:
            payload = json.loads(self.checkpoint_path.read_text(encoding="utf-8"))
        except (ValueError, OSError) as exc:
            log.warning("could not read checkpoint (%s); starting fresh", exc)
            return False

        valid = {f for f in Review.columns()}
        for row in payload.get("reviews", []):
            row = {k: v for k, v in row.items() if k in valid}
            images = row.get("images") or []
            if isinstance(images, str):
                images = [u for u in images.split(" | ") if u]
            row["images"] = images
            review = Review(**row)
            self._reviews[self._key(review)] = review

        self.completed_segments = set(payload.get("completed_segments", []))
        self.total_reviews_reported = payload.get("total_reviews_reported")
        self.total_ratings_reported = payload.get("total_ratings_reported")
        log.info(
            "resumed: %d reviews, %d segments already done",
            len(self._reviews),
            len(self.completed_segments),
        )
        return True

    def clear_checkpoint(self) -> None:
        self.checkpoint_path.unlink(missing_ok=True)

    # -- export -----------------------------------------------------------

    def export(self, formats: list[str]) -> list[Path]:
        """Write each requested format, returning the files actually created.

        One failing format must never cost the user the others: a whole scrape
        can take an hour, and losing the CSV because the spreadsheet writer
        choked on one cell is not an acceptable trade.
        """
        written: list[Path] = []
        for fmt in formats:
            try:
                written.extend(self._export_one(fmt))
            except Exception as exc:
                log.error("could not write %s (other formats unaffected): %s", fmt, exc)
        return written

    def _export_one(self, fmt: str) -> list[Path]:
        """Write a single format. Returns the file written, or [] if skipped."""
        fmt = fmt.strip().lower()
        rows = [r.to_row() for r in self.reviews]
        stem = self.out_dir / f"reviews_{self.asin}"

        if fmt == "csv":
            path = stem.with_suffix(".csv")
            # utf-8-sig so Excel on Windows renders accents correctly.
            with path.open("w", newline="", encoding="utf-8-sig") as fh:
                writer = csv.DictWriter(fh, fieldnames=Review.columns())
                writer.writeheader()
                writer.writerows(rows)
            return [path]

        if fmt == "json":
            path = stem.with_suffix(".json")
            payload = {"summary": self.summary(), "reviews": rows}
            path.write_text(
                json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
            )
            return [path]

        if fmt == "jsonl":
            path = stem.with_suffix(".jsonl")
            with path.open("w", encoding="utf-8") as fh:
                for row in rows:
                    fh.write(json.dumps(row, ensure_ascii=False) + "\n")
            return [path]

        if fmt in ("xlsx", "excel"):
            path = self._export_xlsx(stem.with_suffix(".xlsx"), rows)
            return [path] if path else []

        log.warning("unknown export format %r - skipped", fmt)
        return []

    def _export_xlsx(self, path: Path, rows: list[dict]) -> Path | None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
        except ImportError:
            log.warning("xlsx export needs openpyxl (pip install openpyxl) - skipped")
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "Reviews"
        columns = Review.columns()
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"

        for row in rows:
            ws.append([row.get(c, "") for c in columns])

        widths = {"title": 40, "body": 80, "variant": 25, "permalink": 45, "images": 30}
        for idx, col in enumerate(columns, start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            ws.column_dimensions[letter].width = widths.get(col, 16)
        for row_cells in ws.iter_rows(min_row=2):
            row_cells[columns.index("body")].alignment = Alignment(
                wrap_text=True, vertical="top"
            )

        summary = wb.create_sheet("Summary")
        for key, value in self.summary().items():
            # openpyxl only accepts scalars; dicts and lists must be flattened
            # or it raises "Cannot convert [] to Excel" and kills the export.
            if isinstance(value, (dict, list, tuple, set)):
                value = json.dumps(value if not isinstance(value, (set, tuple)) else list(value))
            summary.append([key, value])
        summary.column_dimensions["A"].width = 30
        summary.column_dimensions["B"].width = 40

        wb.save(path)
        return path
