"""Deduplicating store, checkpointing and exports for Walmart reviews."""

from __future__ import annotations

import csv
import json
import logging
from pathlib import Path

from .models import WalmartPage, WalmartReview

log = logging.getLogger("walmart_reviews.store")


def _flatten(value):
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value if not isinstance(value, (set, tuple)) else list(value))
    return value


class WalmartStore:
    """Collect written reviews for one Walmart item.

    Rating-only entries are dropped at insert time, so every count the UI shows
    means the same thing as the final file.
    """

    def __init__(
        self,
        item_id: str,
        out_dir: Path,
        star_filter: set[int] | None = None,
        cross_product: bool = False,
    ):
        self.item_id = item_id
        self.out_dir = Path(out_dir)
        self.out_dir.mkdir(parents=True, exist_ok=True)
        self.cross_product = cross_product
        self.star_filter = set(star_filter) if star_filter else None
        self._reviews: dict = {}
        self.completed_pages: set[int] = set()
        self.total_reviews: int | None = None
        self.total_ratings: int | None = None
        self.total_pages: int | None = None
        self.average_rating: float | None = None
        self.rating_counts: dict[int, int] = {}
        self.review_rating_counts: dict[int, int] = {}
        self.item_name: str = ""
        self.page_limit_reached: bool = False

    # -- collection -------------------------------------------------------

    def _key(self, review: WalmartReview):
        if self.cross_product:
            return (review.asin, review.review_id)
        return review.review_id

    def _wanted(self, review: WalmartReview) -> bool:
        if not (review.body or "").strip():
            return False
        if self.star_filter is None:
            return True
        if review.rating is None:
            return False
        return int(review.rating) in self.star_filter

    def add_many(self, reviews: list[WalmartReview]) -> int:
        new = 0
        for review in reviews:
            key = self._key(review)
            if key not in self._reviews and self._wanted(review):
                self._reviews[key] = review
                new += 1
        return new

    def capture_page(self, page: WalmartPage) -> None:
        """Record the counts that only the first page reliably reports."""
        if self.total_reviews is None:
            if self.star_filter and page.review_rating_counts:
                self.total_reviews = sum(
                    page.review_rating_counts.get(star, 0) for star in self.star_filter
                )
            else:
                self.total_reviews = page.total_reviews
        if self.total_ratings is None:
            self.total_ratings = page.total_ratings
        if self.total_pages is None:
            self.total_pages = page.total_pages
        if self.average_rating is None:
            self.average_rating = page.average_rating
        if not self.rating_counts:
            self.rating_counts = dict(page.rating_counts)
        if not self.review_rating_counts:
            self.review_rating_counts = dict(page.review_rating_counts)
        if not self.item_name:
            self.item_name = page.item_name

    def __len__(self) -> int:
        return len(self._reviews)

    @property
    def reviews(self) -> list[WalmartReview]:
        return sorted(
            self._reviews.values(),
            key=lambda r: (r.review_date or "0000-00-00"),
            reverse=True,
        )

    # -- summary ----------------------------------------------------------

    @property
    def coverage_denominator(self) -> int | None:
        if self.star_filter and self.review_rating_counts:
            return sum(self.review_rating_counts.get(star, 0) for star in self.star_filter)
        return self.total_reviews

    @property
    def coverage_pct(self) -> float | None:
        denominator = self.coverage_denominator
        if not denominator:
            return None
        return 100.0 * len(self._reviews) / denominator

    def summary(self) -> dict:
        by_star: dict[str, int] = {}
        for review in self._reviews.values():
            key = f"{int(review.rating)}_star" if review.rating else "unrated"
            by_star[key] = by_star.get(key, 0) + 1

        verified = sum(1 for r in self._reviews.values() if r.verified_purchase)
        return {
            "asin": self.item_id,
            "item_name": self.item_name,
            "reviews_downloaded": len(self._reviews),
            "reviews_reported_by_walmart": self.total_reviews,
            "ratings_reported_by_walmart": self.total_ratings,
            "average_rating_reported": self.average_rating,
            "coverage_pct": round(self.coverage_pct, 1) if self.coverage_pct is not None else None,
            "verified_purchases": verified,
            "with_images": sum(1 for r in self._reviews.values() if r.images),
            "by_star": dict(sorted(by_star.items(), reverse=True)),
            "rating_distribution_ratings": {
                str(star): self.rating_counts.get(star, 0) for star in (5, 4, 3, 2, 1)
            },
            "rating_distribution_reviews": {
                str(star): self.review_rating_counts.get(star, 0) for star in (5, 4, 3, 2, 1)
            },
            "pages_completed": len(self.completed_pages),
            "page_limit_reached": self.page_limit_reached,
            "star_filter": sorted(self.star_filter) if self.star_filter else [],
        }

    # -- checkpointing ----------------------------------------------------

    @property
    def checkpoint_path(self) -> Path:
        return self.out_dir / f".checkpoint_walmart_{self.item_id}.json"

    def save_checkpoint(self) -> None:
        payload = {
            "item_id": self.item_id,
            "completed_pages": sorted(self.completed_pages),
            "total_reviews": self.total_reviews,
            "total_ratings": self.total_ratings,
            "total_pages": self.total_pages,
            "average_rating": self.average_rating,
            "rating_counts": {str(k): v for k, v in self.rating_counts.items()},
            "review_rating_counts": {str(k): v for k, v in self.review_rating_counts.items()},
            "item_name": self.item_name,
            "page_limit_reached": self.page_limit_reached,
            "star_filter": sorted(self.star_filter) if self.star_filter else [],
            "reviews": [
                r.to_row() | {"images": r.images}
                for r in self._reviews.values()
            ],
        }
        tmp = self.checkpoint_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        tmp.replace(self.checkpoint_path)

    def load_checkpoint(self) -> bool:
        if not self.checkpoint_path.exists():
            return False
        try:
            payload = json.loads(self.checkpoint_path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            log.warning("could not read Walmart checkpoint (%s); starting fresh", exc)
            return False

        saved_stars = {int(s) for s in (payload.get("star_filter") or [])}
        current_stars = set(self.star_filter or [])
        if saved_stars != current_stars:
            log.info(
                "Walmart checkpoint star filter changed (%s -> %s); starting fresh",
                sorted(saved_stars), sorted(current_stars),
            )
            return False

        valid = set(WalmartReview.columns())
        for row in payload.get("reviews", []):
            if not isinstance(row, dict):
                continue
            row = {k: v for k, v in row.items() if k in valid}
            images = row.get("images") or []
            if isinstance(images, str):
                images = [url for url in images.split(" | ") if url]
            row["images"] = images
            try:
                review = WalmartReview(**row)
            except TypeError:
                continue
            if self._wanted(review):
                self._reviews[self._key(review)] = review

        self.completed_pages = {
            int(p) for p in payload.get("completed_pages", []) if str(p).isdigit()
        }
        self.total_reviews = payload.get("total_reviews")
        self.total_ratings = payload.get("total_ratings")
        self.total_pages = payload.get("total_pages")
        self.average_rating = payload.get("average_rating")
        self.rating_counts = {
            int(k): int(v) for k, v in (payload.get("rating_counts") or {}).items()
        }
        self.review_rating_counts = {
            int(k): int(v) for k, v in (payload.get("review_rating_counts") or {}).items()
        }
        self.item_name = payload.get("item_name") or ""
        self.page_limit_reached = bool(payload.get("page_limit_reached"))
        log.info(
            "resumed Walmart %s: %d reviews, %d pages done",
            self.item_id, len(self._reviews), len(self.completed_pages),
        )
        return True

    def clear_checkpoint(self) -> None:
        self.checkpoint_path.unlink(missing_ok=True)

    # -- exports ----------------------------------------------------------

    def export(self, formats: list[str]) -> list[Path]:
        written: list[Path] = []
        for fmt in formats:
            try:
                written.extend(self._export_one(fmt))
            except Exception as exc:
                log.error("could not write %s (other formats unaffected): %s", fmt, exc)
        return written

    def _export_one(self, fmt: str) -> list[Path]:
        fmt = (fmt or "").strip().lower()
        rows = [r.to_row() for r in self.reviews]
        stem = self.out_dir / f"walmart_reviews_{self.item_id}"

        if fmt == "csv":
            path = stem.with_suffix(".csv")
            with path.open("w", newline="", encoding="utf-8-sig") as fh:
                writer = csv.DictWriter(fh, fieldnames=WalmartReview.columns())
                writer.writeheader()
                writer.writerows(rows)
            return [path]

        if fmt == "json":
            path = stem.with_suffix(".json")
            path.write_text(
                json.dumps(
                    {"summary": self.summary(), "reviews": rows},
                    indent=2,
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            return [path]

        if fmt == "jsonl":
            path = stem.with_suffix(".jsonl")
            with path.open("w", encoding="utf-8") as fh:
                for row in rows:
                    fh.write(json.dumps(row, ensure_ascii=False) + "\n")
            return [path]

        if fmt in ("xlsx", "excel"):
            path = self._export_xlsx(stem.with_suffix(".xlsx"), rows)
            return [path] if path else []

        log.warning("unknown export format %r - skipped", fmt)
        return []

    def _export_xlsx(self, path: Path, rows: list[dict]) -> Path | None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
        except ImportError:
            log.warning("xlsx export needs openpyxl (pip install openpyxl) - skipped")
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "Reviews"
        columns = WalmartReview.columns()
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for row in rows:
            ws.append([row.get(c, "") for c in columns])

        widths = {
            "title": 40, "body": 80, "variant": 25, "permalink": 45,
            "images": 30, "seller": 22, "source": 16,
        }
        for idx, col in enumerate(columns, start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            ws.column_dimensions[letter].width = widths.get(col, 16)
        for row_cells in ws.iter_rows(min_row=2):
            row_cells[columns.index("body")].alignment = Alignment(
                wrap_text=True, vertical="top"
            )

        summary = wb.create_sheet("Summary")
        for key, value in self.summary().items():
            summary.append([key, _flatten(value)])
        summary.column_dimensions["A"].width = 32
        summary.column_dimensions["B"].width = 48

        ratings = wb.create_sheet("Ratings")
        ratings.append(["stars", "ratings", "written_reviews"])
        for cell in ratings[1]:
            cell.font = Font(bold=True)
        for star in (5, 4, 3, 2, 1):
            ratings.append([
                star,
                self.rating_counts.get(star, 0),
                self.review_rating_counts.get(star, 0),
            ])
        for letter, width in (("A", 10), ("B", 14), ("C", 18)):
            ratings.column_dimensions[letter].width = width

        wb.save(path)
        return path


def export_combined(
    stores: list[WalmartStore],
    out_dir: Path,
    formats: list[str],
    stem: str = "all_products",
) -> list[Path]:
    """Merge every product into one file, keeping one row per (item, review)."""
    stores = [store for store in stores if len(store)]
    if not stores:
        return []

    merged = WalmartStore(stem, out_dir, cross_product=True)
    for store in stores:
        merged.add_many(store.reviews)

    totals = [s.total_reviews for s in stores if s.total_reviews]
    if totals:
        merged.total_reviews = sum(totals)
    ratings = [s.total_ratings for s in stores if s.total_ratings]
    if ratings:
        merged.total_ratings = sum(ratings)

    for star in (5, 4, 3, 2, 1):
        merged.rating_counts[star] = sum(s.rating_counts.get(star, 0) for s in stores)
        merged.review_rating_counts[star] = sum(
            s.review_rating_counts.get(star, 0) for s in stores
        )

    weighted = [
        (s.average_rating, s.total_ratings)
        for s in stores
        if s.average_rating is not None and s.total_ratings
    ]
    if weighted:
        merged.average_rating = round(
            sum(avg * count for avg, count in weighted)
            / sum(count for _, count in weighted),
            2,
        )

    written = merged.export(formats)
    merged.clear_checkpoint()
    return written
