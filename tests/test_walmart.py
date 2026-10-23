"""Offline tests for the independent Walmart implementation.

Run with:  python tests/test_walmart.py
       or: python -m pytest tests -q
"""

from __future__ import annotations

import argparse
import ast
import csv
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from walmart_reviews.models import (  # noqa: E402
    WalmartReview,
    extract_item_ids,
    parse_item_id,
)
from walmart_reviews.parser import (  # noqa: E402
    detect_block,
    parse_next_data,
)
from walmart_reviews.cli import parse_stars  # noqa: E402
from walmart_reviews.store import WalmartStore, export_combined  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).parent / "fixtures"
PAGE1 = FIXTURES / "walmart_reviews_page1.json"
PAGE2 = FIXTURES / "walmart_reviews_page2.json"


def _payload(path: Path = PAGE1) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _html(payload: dict) -> str:
    blob = json.dumps({
        "props": {
            "pageProps": {
                "initialData": {"data": {"reviews": payload}},
            },
        },
    })
    return (
        "<html><head></head><body>"
        f'<script id="__NEXT_DATA__" type="application/json">{blob}</script>'
        "</body></html>"
    )


def _page(path: Path = PAGE1, number: int = 1):
    return parse_next_data(
        _html(_payload(path)),
        "20951672018",
        page=number,
        found_via=f"walmart/relevancy/p{number}",
    )


# -- item IDs -------------------------------------------------------------


def test_item_id_accepts_walmart_forms():
    assert parse_item_id("20951672018") == "20951672018"
    assert parse_item_id("20951672018/") == "20951672018"
    assert parse_item_id("https://www.walmart.com/ip/Apple-AirPods-5/20951672018") == "20951672018"
    assert parse_item_id("https://www.walmart.com/ip/20951672018") == "20951672018"
    assert parse_item_id(
        "https://www.walmart.com/reviews/product/20951672018?sort=relevancy&page=2"
    ) == "20951672018"


def test_item_id_rejects_amazon_asin_and_short_numbers():
    assert parse_item_id("B08H2Y91P1") is None
    assert parse_item_id("12345") is None
    assert parse_item_id("https://www.amazon.com/dp/B08H2Y91P1") is None


def test_extract_item_ids_preserves_order_and_dedupes():
    raw = """
    20951672018
    https://www.walmart.com/ip/123456789
    20951672018
    not-an-id
    """
    assert extract_item_ids(raw) == ["20951672018", "123456789"]


# -- parsing --------------------------------------------------------------


def test_parse_page_keeps_only_written_reviews():
    page = _page()
    assert len(page.reviews) == 2, "the whitespace-only review must be filtered"
    assert [r.review_id for r in page.reviews] == ["100000001", "100000002"]


def test_parse_page_fields():
    first = _page().reviews[0]
    assert first.title == "Great purchase."
    assert first.body == "Works well and arrived quickly."
    assert first.rating == 5.0
    assert first.author == "TestUser"
    assert first.review_date == "2026-10-13"
    assert first.country == "Texas"
    assert first.verified_purchase is True
    assert first.helpful_votes == 2
    assert first.images == ["https://example.com/review-photo.jpg"]
    assert first.seller == "Walmart.com"
    assert first.source == "bazaarvoice"
    assert first.permalink.endswith("#review-100000001")
    assert first.asin == "20951672018"
    assert first.found_via == "walmart/relevancy/p1"


def test_null_title_is_kept_not_dropped():
    second = _page().reviews[1]
    assert second.title == ""
    assert second.body == "Good value for the price."
    assert second.verified_purchase is False


def test_counts_average_and_distribution():
    page = _page()
    assert page.total_reviews == 21
    assert page.total_ratings == 273
    assert page.average_rating == 4.5
    assert page.rating_counts == {5: 212, 4: 24, 3: 12, 2: 2, 1: 23}
    assert page.review_rating_counts == {5: 12, 4: 2, 3: 1, 2: 1, 1: 5}


def test_pagination_uses_pages_not_next():
    page = _page(PAGE2, number=2)
    assert page.current_span == "11-20"
    assert page.total_pages == 3
    assert page.has_next is True
    assert page.reviews[0].review_id == "100000011"


def test_challenge_and_notfound_detection():
    assert detect_block(
        "<html><body>Robot or human? Activate and hold</body></html>",
        "https://www.walmart.com/blocked?url=x",
    ) == "challenge"
    assert detect_block(
        "<html><head><title>Page Not Found</title></head></html>",
        "https://www.walmart.com/ip/0",
    ) == "notfound"
    assert detect_block(_html(_payload())) is None


def test_missing_next_data_raises():
    from walmart_reviews.models import WalmartParseError

    try:
        parse_next_data("<html><body>nothing</body></html>", "20951672018")
    except WalmartParseError:
        return
    raise AssertionError("parse_next_data should raise when the payload is missing")


# -- store / checkpoint / export -----------------------------------------


def test_store_filters_empty_body_and_star():
    with tempfile.TemporaryDirectory() as tmp:
        store = WalmartStore("20951672018", Path(tmp), star_filter={5})
        empty = WalmartReview("empty", body="   ", rating=5)
        five = WalmartReview("five", body="text", rating=5)
        four = WalmartReview("four", body="text", rating=4)
        added = store.add_many([empty, five, four])
        assert added == 1
        assert [r.review_id for r in store.reviews] == ["five"]


def test_store_summary_and_coverage():
    with tempfile.TemporaryDirectory() as tmp:
        store = WalmartStore("20951672018", Path(tmp))
        page = _page()
        store.capture_page(page)
        store.add_many(page.reviews)
        summary = store.summary()
        assert summary["reviews_downloaded"] == 2
        assert summary["reviews_reported_by_walmart"] == 21
        assert summary["ratings_reported_by_walmart"] == 273
        assert summary["average_rating_reported"] == 4.5
        assert summary["coverage_pct"] == 9.5
        assert summary["rating_distribution_reviews"]["5"] == 12


def test_checkpoint_roundtrip():
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        store = WalmartStore("20951672018", out)
        page = _page()
        store.capture_page(page)
        store.add_many(page.reviews)
        store.completed_pages.add(1)
        store.save_checkpoint()

        again = WalmartStore("20951672018", out)
        assert again.load_checkpoint() is True
        assert len(again) == 2
        assert again.completed_pages == {1}
        assert again.total_reviews == 21
        assert again.total_pages == 3
        assert again.total_ratings == 273


def test_all_exports_have_walmart_columns_and_sheets():
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        store = WalmartStore("20951672018", out)
        page = _page()
        store.capture_page(page)
        store.add_many(page.reviews)
        written = store.export(["csv", "json", "xlsx"])
        suffixes = {p.suffix for p in written}
        assert suffixes == {".csv", ".json", ".xlsx"}, suffixes

        csv_path = next(p for p in written if p.suffix == ".csv")
        with csv_path.open(encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
        assert len(rows) == 2
        assert list(rows[0].keys()) == WalmartReview.columns()
        assert "seller" in rows[0] and "source" in rows[0]
        assert "record_type" not in rows[0]

        from openpyxl import load_workbook

        xlsx = next(p for p in written if p.suffix == ".xlsx")
        wb = load_workbook(xlsx)
        assert {"Reviews", "Summary", "Ratings"} <= set(wb.sheetnames)


def test_combined_export_keeps_one_row_per_item():
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        a = WalmartStore("111111111", out)
        b = WalmartStore("222222222", out)
        page = _page()
        a.capture_page(page)
        b.capture_page(page)
        for store in (a, b):
            store.add_many([
                WalmartReview(**{**r.to_row(), "images": r.images, "asin": store.item_id})
                for r in page.reviews
            ])

        written = export_combined([a, b], out, ["csv"])
        with written[0].open(encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
        assert len(rows) == 4
        assert {r["asin"] for r in rows} == {"111111111", "222222222"}


# -- CLI / routes / isolation --------------------------------------------


def test_parse_stars():
    assert parse_stars("all") is None
    assert parse_stars("5,4") == {5, 4}
    assert parse_stars("") is None
    for bad in ("0", "6", "5,x"):
        try:
            parse_stars(bad)
        except argparse.ArgumentTypeError:
            continue
        raise AssertionError(f"{bad} should be rejected")


def test_walmart_page_and_switch_routes():
    from webapp.server import app

    client = app.test_client()
    amazon = client.get("/")
    assert amazon.status_code == 200
    assert b"/walmart" in amazon.data

    walmart = client.get("/walmart")
    assert walmart.status_code == 200
    assert b"/static/walmart.js" in walmart.data
    assert b'aria-current="page">Walmart' in walmart.data

    status = client.get("/walmart/api/status")
    assert status.status_code == 200
    assert status.get_json()["running"] is False

    # Amazon status must remain intact.
    assert client.get("/api/status").status_code == 200

    # No path traversal through the Walmart download route.
    assert client.get("/walmart/api/download/..%2F..%2Fetc%2Fpasswd").status_code == 404


def test_walmart_code_does_not_import_amazon():
    files = list((ROOT / "walmart_reviews").glob("*.py"))
    files.append(ROOT / "webapp" / "walmart_server.py")
    for path in files:
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                names = [alias.name for alias in node.names]
            elif isinstance(node, ast.ImportFrom):
                names = [node.module or ""]
            else:
                continue
            for name in names:
                assert not name.startswith("amazon_reviews"), f"{path} imports {name}"


def test_star_filter_uses_matching_written_review_denominator():
    with tempfile.TemporaryDirectory() as tmp:
        store = WalmartStore("20951672018", Path(tmp), star_filter={5})
        page = _page()
        store.capture_page(page)
        store.add_many(page.reviews)
        summary = store.summary()
        assert summary["reviews_reported_by_walmart"] == 12
        assert summary["coverage_pct"] == 8.3


def test_checkpoint_is_ignored_when_star_filter_changes():
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        all_stars = WalmartStore("20951672018", out)
        page = _page()
        all_stars.capture_page(page)
        all_stars.add_many(page.reviews)
        all_stars.completed_pages.add(1)
        all_stars.save_checkpoint()

        five_only = WalmartStore("20951672018", out, star_filter={5})
        assert five_only.load_checkpoint() is False
        assert len(five_only) == 0


def test_start_endpoint_runs_worker_with_fake_scraper():
    """Exercise the Flask job/worker path without touching Walmart."""
    from unittest.mock import patch
    import webapp.walmart_server as ws

    page = _page()

    class FakeScraper:
        def __init__(self, item_id, store, **kwargs):
            self.store = store
            self.on_event = kwargs.get("on_event") or (lambda **kw: None)

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def retarget(self, item_id, store):
            self.store = store

        def run(self):
            self.store.capture_page(page)
            self.store.add_many(page.reviews)
            self.store.completed_pages.add(1)
            self.store.total_pages = 3
            self.store.save_checkpoint()
            self.on_event(
                stage="scraping", page=1, total_pages=3,
                reviews=len(self.store), reviews_total=self.store.total_reviews,
            )
            return self.store

    with tempfile.TemporaryDirectory() as tmp, patch.object(ws, "WalmartScraper", FakeScraper), \
            patch.object(ws, "OUTPUT_DIR", Path(tmp)):
        from webapp.server import app
        client = app.test_client()
        response = client.post("/walmart/api/start", json={"items": "20951672018"})
        assert response.status_code == 200, response.get_data(as_text=True)
        ws.JOB.thread.join(10)
        assert ws.JOB.running is False
        assert len(ws.JOB.results) == 1
        entry = ws.JOB.results[0]
        assert entry["ok"] is True
        assert entry["reviews"] == 2
        assert entry["files"]
        assert (Path(tmp) / entry["files"][0]).exists()


if __name__ == "__main__":
    failed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"  PASS  {name}")
            except AssertionError as exc:
                failed += 1
                print(f"  FAIL  {name}: {exc}")
            except Exception as exc:  # noqa: BLE001
                failed += 1
                print(f"  ERROR {name}: {type(exc).__name__}: {exc}")
    print(f"\n{'FAILED' if failed else 'ALL PASSED'} ({failed} failures)")
    sys.exit(1 if failed else 0)
